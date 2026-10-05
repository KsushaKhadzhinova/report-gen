from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt
from PIL import Image

from reportgen.domain import enterprise_standard as standard
from reportgen.domain.blocks import Block, Kind
from reportgen.domain.title_page import CENTER, FIELD, SIGNATURE, build_title_page
from reportgen.infrastructure.safe_paths import resolve_inside

INLINE_RE = re.compile(r"(`[^`]+`|\*[^*]+\*)")
FIGURE_MAX_WIDTH_CM = 15.5
FIGURE_MAX_HEIGHT_CM = 20.0
TITLE_FIELD_TAB_CM = 3.5


def _configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = standard.FONT
    normal.font.size = Pt(standard.FONT_SIZE_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), standard.FONT)
    fmt = normal.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(standard.LINE_SPACING_PT)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = standard.FONT
        style.font.size = Pt(standard.FONT_SIZE_PT)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = None
        style.element.rPr.rFonts.set(qn("w:eastAsia"), standard.FONT)
        style.element.rPr.rFonts.set(qn("w:ascii"), standard.FONT)
        style.element.rPr.rFonts.set(qn("w:hAnsi"), standard.FONT)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(standard.LINE_SPACING_PT)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(standard.LINE_SPACING_PT)
        style.paragraph_format.keep_with_next = True


def _configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = Mm(standard.MARGIN_LEFT_MM)
    section.right_margin = Mm(standard.MARGIN_RIGHT_MM)
    section.top_margin = Mm(standard.MARGIN_TOP_MM)
    section.bottom_margin = Mm(standard.MARGIN_BOTTOM_MM)


def _field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instruction), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        elif text == instruction:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = instruction
        else:
            element = OxmlElement("w:t")
            element.text = text
        run._r.append(element)


def _page_number_footer(doc: Document) -> None:
    footer = doc.sections[0].footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _field(paragraph, "PAGE")


def _add_runs(paragraph, text: str) -> None:
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(12)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _body(doc: Document, text: str, indent: bool = True):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        paragraph.paragraph_format.first_line_indent = Cm(standard.PARAGRAPH_INDENT_CM)
    _add_runs(paragraph, text)
    return paragraph


def _blank(doc: Document) -> None:
    doc.add_paragraph()


def _page_break_unless_at_start(doc: Document) -> None:
    if doc.paragraphs and doc.paragraphs[-1].text:
        doc.add_page_break()


def _centered_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Heading 1")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(text)


def _heading(doc: Document, block: Block) -> None:
    if block.appendix:
        _page_break_unless_at_start(doc)
        _centered_heading(doc, f"ПРИЛОЖЕНИЕ {block.number}")
        _centered_heading(doc, block.text)
        return
    if block.level == 1:
        _page_break_unless_at_start(doc)
    paragraph = doc.add_paragraph(style=f"Heading {block.level}")
    label = f"{block.number} " if block.number else ""
    paragraph.add_run(f"{label}{block.text}").bold = True
    if block.is_centered_heading:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Cm(standard.PARAGRAPH_INDENT_CM)


def _picture_width_cm(path: Path) -> float:
    with Image.open(path) as image:
        ratio = image.height / image.width
    return min(FIGURE_MAX_WIDTH_CM, FIGURE_MAX_HEIGHT_CM / ratio)


def _figure(doc: Document, block: Block, base: Path) -> None:
    path = resolve_inside(base, block.path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    if path:
        paragraph.add_run().add_picture(str(path), width=Cm(_picture_width_cm(path)))
    else:
        paragraph.add_run(f"[нет файла: {block.path}]")
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.add_run(standard.figure_caption(block.number, block.caption))
    _blank(doc)


def _set_cell_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    properties.append(borders)


def _table(doc: Document, block: Block) -> None:
    if block.caption:
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
        caption.paragraph_format.keep_with_next = True
        caption.add_run(standard.table_caption(block.number, block.caption))
    rows = block.rows
    columns = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_cell_borders(table)
    for r, row in enumerate(rows):
        for c in range(columns):
            cell = table.cell(r, c)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if r == 0 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            run = paragraph.add_run(row[c] if c < len(row) else "")
            run.font.size = Pt(12)
    _blank(doc)


def _code(doc: Document, block: Block) -> None:
    if block.caption:
        caption = doc.add_paragraph()
        caption.paragraph_format.keep_with_next = True
        caption.add_run(standard.listing_caption(block.number, block.caption))
    for line in block.text.splitlines() or [""]:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = paragraph.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(11)
    _blank(doc)


def _title_page(doc: Document, meta: dict) -> None:
    text_width = Mm(210 - standard.MARGIN_LEFT_MM - standard.MARGIN_RIGHT_MM)
    for line in build_title_page(meta):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        if line.kind == CENTER:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(line.text).bold = line.bold
        elif line.kind == FIELD:
            paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(TITLE_FIELD_TAB_CM))
            paragraph.add_run(f"{line.label}\t{line.text}" if not line.label.endswith(":") else f"{line.label} {line.text}")
        elif line.kind == SIGNATURE:
            paragraph.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
            paragraph.add_run(f"{line.label}\t{line.text}")
    doc.add_page_break()


def _toc(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("СОДЕРЖАНИЕ").bold = True
    _blank(doc)
    paragraph = doc.add_paragraph()
    _field(paragraph, 'TOC \\o "1-2" \\h \\z \\u')
    doc.add_page_break()


def _request_field_update(doc: Document) -> None:
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def _list(doc: Document, block: Block) -> None:
    for line in standard.list_items(block.items):
        _body(doc, line)


class DocxRenderer:
    def render(self, blocks: list[Block], meta: dict, base_dir: Path, output: Path) -> Path:
        doc = Document()
        _configure_styles(doc)
        _configure_page(doc)
        _page_number_footer(doc)
        if meta.get("title_page", True):
            _title_page(doc, meta)
        if meta.get("toc", True):
            _toc(doc)
        for block in blocks:
            self._render_block(doc, block, base_dir)
        _request_field_update(doc)
        output.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output)
        return output

    @staticmethod
    def _render_block(doc: Document, block: Block, base_dir: Path) -> None:
        handlers = {
            Kind.HEADING: lambda: _heading(doc, block),
            Kind.PARAGRAPH: lambda: _body(doc, block.text),
            Kind.LIST: lambda: _list(doc, block),
            Kind.FIGURE: lambda: _figure(doc, block, base_dir),
            Kind.TABLE: lambda: _table(doc, block),
            Kind.CODE: lambda: _code(doc, block),
        }
        handlers[block.kind]()
