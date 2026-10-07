from __future__ import annotations

import copy
import hashlib
import re
from dataclasses import dataclass, replace

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from reportgen.domain import enterprise_standard as standard
from reportgen.domain.citations import CITATION_RE, parse_numbers, renumber
from reportgen.domain.fix_options import FixOptions
from reportgen.domain.object_sentences import (
    FIGURE_RE,
    FIGURE_WORDS,
    TABLE_RE,
    TABLE_WORDS,
    figure_reference_sentence,
    reference_number_pattern,
    table_reference_sentence,
)
from reportgen.domain.sources import vak_entries

BLANK_LINE = Pt(standard.LINE_SPACING_PT)
TABLE_FONT_PT = 12
KEEP_TOGETHER_MAX_ROWS = 12
TEXT_WIDTH_CM = 21.0 - standard.MARGIN_LEFT_MM / 10 - standard.MARGIN_RIGHT_MM / 10
CHAR_WIDTH_CM = 0.25
CELL_PADDING_CM = 0.45
BREAKABLE_WORD_CHARS = 22
MONTHS = "январь|февраль|март|апрель|май|июнь|июль|август|сентябрь|октябрь|ноябрь|декабрь"
CITY_DATE_RE = re.compile(rf"^(Минск),?\s+(?:{MONTHS})\s+(\d{{4}})\s*$", re.IGNORECASE)
TABLE_CAPTION_RE = re.compile(r"^Таблица\s+[\dА-Я]+(\.\d+)?\s*[–-]")
FIGURE_CAPTION_RE = re.compile(r"^Рисунок\s+[\dА-Я]+(\.\d+)?\s*[–-]")
LEADING_MARK_RE = re.compile(r"^\s*[–—\-•]\s*")
FILLER_PHRASES = ("В данном подразделе рассматриваются ключевые аспекты",)
HEADING_ENDS = ("Heading",)
SOURCES_TITLE = "список использованных источников"
APPENDIX_RE = re.compile(r"^ПРИЛОЖЕНИЕ\s+[А-ЯA-Z]\.?$", re.IGNORECASE)
APPENDIX_STATUS_RE = re.compile(r"^\((обязательное|справочное|рекомендуемое)\)$", re.IGNORECASE)
APPENDIX_TITLE_MAX_CHARS = 160


def _paragraphs_and_tables(document):
    """Элементы тела документа по порядку: абзацы и таблицы."""
    for element in document.element.body.iterchildren():
        if element.tag == qn("w:p"):
            yield Paragraph(element, document)
        elif element.tag == qn("w:tbl"):
            yield Table(element, document)


def _has_picture(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing"))


def _is_heading(paragraph) -> bool:
    return isinstance(paragraph, Paragraph) and paragraph.style.name.startswith(HEADING_ENDS) and bool(paragraph.text.strip())


def _set_text(paragraph, text: str) -> None:
    for run in paragraph.runs[1:]:
        run._r.getparent().remove(run._r)
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def fix_title_date(document) -> int:
    """На титульном листе «Минск 2026» без месяца."""
    fixed = 0
    for paragraph in document.paragraphs[:60]:
        match = CITY_DATE_RE.match(paragraph.text.strip())
        if match:
            _set_text(paragraph, f"{match.group(1)} {match.group(2)}")
            fixed += 1
    return fixed


def fix_section_headings(document) -> int:
    """Заголовки разделов набираются прописными буквами и в тексте, и в содержании."""
    fixed = 0
    for paragraph in document.paragraphs:
        if paragraph.style.name != "Heading 1" or not paragraph.text.strip():
            continue
        if paragraph.text != paragraph.text.upper():
            _set_text(paragraph, paragraph.text.upper())
            fixed += 1
    return fixed


def remove_filler_paragraphs(document) -> int:
    """Убирает шаблонные фразы-заглушки, не несущие смысла."""
    removed = 0
    for paragraph in list(document.paragraphs):
        if paragraph.text.strip().startswith(FILLER_PHRASES):
            paragraph._p.getparent().remove(paragraph._p)
            removed += 1
    return removed


def _space_before(paragraph, value) -> None:
    paragraph.paragraph_format.space_before = value


def fix_subsection_spacing(document) -> int:
    fixed = 0
    for paragraph in document.paragraphs:
        if paragraph.style.name in ("Heading 2", "Heading 3") and paragraph.text.strip():
            _space_before(paragraph, BLANK_LINE)
            fixed += 1
    return fixed


def _fix_figure(paragraph, caption: Paragraph | None) -> None:
    fmt = paragraph.paragraph_format
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.first_line_indent = Cm(0)
    fmt.space_before = BLANK_LINE
    fmt.space_after = Pt(0)
    fmt.keep_with_next = True
    if caption is not None:
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.first_line_indent = Cm(0)
        caption.paragraph_format.space_before = Pt(0)
        caption.paragraph_format.space_after = BLANK_LINE


def _fix_table_cell_paragraphs(table: Table) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                fmt = paragraph.paragraph_format
                fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
                fmt.first_line_indent = Cm(0)
                fmt.space_before = Pt(0)
                fmt.space_after = Pt(0)
                if paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.name = standard.FONT
                    run.font.size = Pt(TABLE_FONT_PT)


def _has_merged_cells(table: Table) -> bool:
    return bool(table._tbl.xpath(".//w:gridSpan | .//w:vMerge"))


def _column_texts(table: Table) -> list[list[str]]:
    columns: list[list[str]] = [[] for _ in table.columns]
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            columns[index].append(cell.text)
    return columns


def _shrink_wide_columns(minimum: list[float], total: float) -> list[float]:
    """Не хватает места: узкие столбцы (номера, даты, короткие слова) остаются целыми, сжимаются только широкие."""
    widths = [0.0] * len(minimum)
    remaining = list(range(len(minimum)))
    space = total
    while remaining:
        share = space / len(remaining)
        small = [index for index in remaining if minimum[index] <= share]
        if not small:
            scale = space / sum(minimum[index] for index in remaining)
            for index in remaining:
                widths[index] = minimum[index] * scale
            break
        for index in small:
            widths[index] = minimum[index]
            space -= minimum[index]
            remaining.remove(index)
    return widths


def column_widths_cm(columns: list[list[str]], total: float = TEXT_WIDTH_CM) -> list[float]:
    """Ширины столбцов: минимум по самому длинному слову, остаток пропорционально объёму текста; очень длинные слова допускается переносить."""
    minimum, preferred = [], []
    for cells in columns:
        longest = max((len(word) for text in cells for word in text.split()), default=1)
        widest = max((len(text) for text in cells), default=1)
        low = min(longest, BREAKABLE_WORD_CHARS) * CHAR_WIDTH_CM + CELL_PADDING_CM
        minimum.append(low)
        preferred.append(max(low, widest * CHAR_WIDTH_CM + CELL_PADDING_CM))
    if sum(preferred) <= total:
        spare = total - sum(preferred)
        return [width + spare * width / sum(preferred) for width in preferred]
    if sum(minimum) >= total:
        return _shrink_wide_columns(minimum, total)
    room = total - sum(minimum)
    extra = [high - low for high, low in zip(preferred, minimum, strict=True)]
    return [low + room * share / sum(extra) for low, share in zip(minimum, extra, strict=True)]


def _make_columns_fit_content(table: Table) -> None:
    if _has_merged_cells(table):
        table.autofit = True
        return
    widths = column_widths_cm(_column_texts(table))
    table.autofit = False
    properties = table._tbl.tblPr
    total = properties.find(qn("w:tblW"))
    if total is None:
        total = OxmlElement("w:tblW")
        properties.append(total)
    total.set(qn("w:type"), "dxa")
    total.set(qn("w:w"), str(int(sum(widths) * 567)))
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Cm(width)
        for cell in column.cells:
            cell.width = Cm(width)


def _mark_header_and_keep_rows(table: Table) -> None:
    rows = table.rows
    for index, row in enumerate(rows):
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn("w:cantSplit")) is None:
            properties.append(OxmlElement("w:cantSplit"))
        if index == 0 and properties.find(qn("w:tblHeader")) is None:
            properties.append(OxmlElement("w:tblHeader"))
    if len(rows) <= KEEP_TOGETHER_MAX_ROWS:
        for row in rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True


def fix_objects(document) -> int:
    """Пустые строки и интервалы вокруг рисунков и таблиц, вид таблиц."""
    fixed = 0
    items = list(_paragraphs_and_tables(document))
    for index, item in enumerate(items):
        following = items[index + 1] if index + 1 < len(items) else None
        if isinstance(item, Table):
            _fix_table_cell_paragraphs(item)
            _make_columns_fit_content(item)
            _mark_header_and_keep_rows(item)
            if isinstance(following, Paragraph) and following.text.strip() and not TABLE_CAPTION_RE.match(following.text.strip()):
                _space_before(following, BLANK_LINE)
            fixed += 1
        elif TABLE_CAPTION_RE.match(item.text.strip()):
            fmt = item.paragraph_format
            fmt.space_before = BLANK_LINE
            fmt.space_after = Pt(0)
            fmt.keep_with_next = True
            fixed += 1
        elif _has_picture(item) and len(item.text.strip()) < 3:
            caption = following if isinstance(following, Paragraph) and FIGURE_CAPTION_RE.match(following.text.strip()) else None
            _fix_figure(item, caption)
            fixed += 1
    return fixed


def _is_blank(item) -> bool:
    return isinstance(item, Paragraph) and not item.text.strip() and not _has_picture(item)


CONTINUATION_RE = re.compile(r"^Продолжение таблицы\s")


def _is_continuation(item) -> bool:
    """«Продолжение таблицы N» между частями длинной таблицы, которые Word делит по страницам."""
    return isinstance(item, Paragraph) and CONTINUATION_RE.match(item.text.strip()) is not None


def _is_caption(item) -> bool:
    return isinstance(item, Paragraph) and bool(TABLE_CAPTION_RE.match(item.text.strip()) or FIGURE_CAPTION_RE.match(item.text.strip()))


def _is_plain_text(item) -> bool:
    """Обычный абзац текста: не заголовок, не подпись, не рисунок и не пустая строка."""
    return (
        isinstance(item, Paragraph)
        and bool(item.text.strip())
        and not _has_picture(item)
        and not _is_heading(item)
        and not _is_caption(item)
        and CONTINUATION_RE.match(item.text.strip()) is None
    )


@dataclass(frozen=True)
class _Object:
    """Таблица или рисунок в списке содержимого без пустых строк: индексы первого и последнего элемента объекта."""

    kind: str
    number: str
    caption: str
    first: int
    last: int
    table: Table | None = None
    picture: Paragraph | None = None


def _content_items(document) -> list:
    return [item for item in _paragraphs_and_tables(document) if not _is_blank(item)]


def _find_objects(items: list) -> list[_Object]:
    found = []
    for index, item in enumerate(items):
        if not isinstance(item, Paragraph):
            continue
        text = item.text.strip()
        table_match = TABLE_RE.match(text) if not CONTINUATION_RE.match(text) else None
        figure_match = FIGURE_RE.match(text)
        if table_match and index + 1 < len(items) and isinstance(items[index + 1], Table):
            last = index + 1
            while last + 2 < len(items) and _is_continuation(items[last + 1]) and isinstance(items[last + 2], Table):
                last += 2
            found.append(_Object("table", table_match["number"], text, index, last, table=items[index + 1]))
        elif figure_match:
            before = items[index - 1] if index else None
            picture = before if isinstance(before, Paragraph) and _has_picture(before) else None
            found.append(_Object("figure", figure_match["number"], text, index - 1 if picture is not None else index, index, picture=picture))
    return found


def _has_reference(paragraph: Paragraph, obj: _Object) -> bool:
    words = TABLE_WORDS if obj.kind == "table" else FIGURE_WORDS
    return reference_number_pattern(obj.number, words).search(paragraph.text) is not None


def _reference_state(items: list, obj: _Object) -> str:
    """Перед объектом должен стоять текст со ссылкой на него: «ok», «colon» (абзац оканчивается двоеточием) или «missing»."""
    previous = items[obj.first - 1] if obj.first else None
    if not _is_plain_text(previous):
        return "missing"
    if _has_reference(previous, obj):
        return "ok"
    return "colon" if previous.text.rstrip().endswith(":") else "missing"


def _has_text_after(items: list, obj: _Object) -> bool:
    following = items[obj.last + 1] if obj.last + 1 < len(items) else None
    return _is_plain_text(following)


def count_objects_without_text(document) -> int:
    """Таблицы и рисунки, после которых нет абзаца текста."""
    items = _content_items(document)
    return sum(not _has_text_after(items, obj) for obj in _find_objects(items))


def count_objects_without_reference(document) -> int:
    """Таблицы и рисунки, перед которыми нет текста со ссылкой на них."""
    items = _content_items(document)
    return sum(_reference_state(items, obj) != "ok" for obj in _find_objects(items))


def _table_text(table: Table) -> str:
    return "\x1f".join(cell.text for row in table.rows for cell in row.cells)


def _image_blob(document, picture: Paragraph | None) -> bytes | None:
    if picture is None:
        return None
    ids = picture._p.xpath(".//a:blip/@r:embed")
    return document.part.related_parts[ids[0]].blob if ids else None


def _object_key(document, obj: _Object) -> str:
    """Ключ абзаца для объекта: подпись и отпечаток содержимого, чтобы одинаковые подписи разных объектов не путались."""
    payload = _table_text(obj.table).encode() if obj.table is not None else (_image_blob(document, obj.picture) or obj.caption.encode())
    return f"{obj.caption}|{hashlib.sha1(payload).hexdigest()[:10]}"


def _heading_before(items: list, index: int) -> str:
    for earlier in reversed(items[:index]):
        if _is_heading(earlier):
            return earlier.text.strip()
    return ""


def _variants(objects: list[_Object]) -> dict[int, int]:
    """Номер формулировки предложения со ссылкой: чередуется по порядку объектов одного вида."""
    counters = {"table": 0, "figure": 0}
    result = {}
    for obj in objects:
        result[obj.first] = counters[obj.kind]
        counters[obj.kind] += 1
    return result


def _reference_sentence(obj: _Object, variant: int) -> str | None:
    return (table_reference_sentence if obj.kind == "table" else figure_reference_sentence)(obj.caption, variant)


def objects_needing_notes(document, cell_chars: int = 140, max_rows: int = 30) -> list[dict]:
    """Объекты, после которых нужно дописать абзац: подпись, данные таблицы, окружающий текст; абзац пишется автором по этим данным."""
    items = _content_items(document)
    objects = _find_objects(items)
    variants = _variants(objects)
    needed = []
    for obj in objects:
        if _has_text_after(items, obj):
            continue
        previous = items[obj.first - 1] if obj.first else None
        following = items[obj.last + 1] if obj.last + 1 < len(items) else None
        entry = {
            "key": _object_key(document, obj),
            "kind": obj.kind,
            "caption": obj.caption,
            "section": _heading_before(items, obj.first),
            "text_before": previous.text.strip()[:400] if _is_plain_text(previous) else "",
            "reference_sentence": _reference_sentence(obj, variants[obj.first]) if _reference_state(items, obj) != "ok" else "",
            "next_heading": following.text.strip()[:120] if _is_heading(following) else "",
            "text": "",
        }
        if obj.table is not None:
            rows = [[" ".join(cell.text.split())[:cell_chars] for cell in row.cells] for row in obj.table.rows]
            entry["rows"] = rows[:max_rows]
            entry["total_rows"] = len(rows)
        needed.append(entry)
    return needed


def figure_image(document, key: str) -> tuple[str, bytes] | None:
    """Файл рисунка по ключу абзаца, если ключ относится к рисунку."""
    items = _content_items(document)
    for obj in _find_objects(items):
        if obj.kind == "figure" and _object_key(document, obj) == key:
            blob = _image_blob(document, obj.picture)
            return (f"{hashlib.sha1(blob).hexdigest()[:10]}.png", blob) if blob else None
    return None


def _new_paragraph(document, text: str, space_before: Pt = Pt(0)) -> Paragraph:
    paragraph = Paragraph(OxmlElement("w:p"), document._body)
    run = paragraph.add_run(text)
    run.font.name = standard.FONT
    run.font.size = Pt(standard.FONT_SIZE_PT)
    fmt = paragraph.paragraph_format
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.first_line_indent = Cm(standard.PARAGRAPH_INDENT_CM)
    fmt.left_indent = Cm(0)
    fmt.right_indent = Cm(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(standard.LINE_SPACING_PT)
    fmt.space_before = space_before
    fmt.space_after = Pt(0)
    return paragraph


def _element(item):
    return item._tbl if isinstance(item, Table) else item._p


def _add_reference_before_colon(paragraph: Paragraph, obj: _Object) -> None:
    runs = [run for run in paragraph.runs if run.text]
    last = runs[-1]
    word = "таблица" if obj.kind == "table" else "рисунок"
    last.text = f"{last.text.rstrip()[:-1].rstrip()} ({word} {obj.number}):"


def fix_text_around_objects(document, options: FixOptions = FixOptions()) -> int:
    """Перед таблицей или рисунком ставится текст со ссылкой на них, после них абзац текста по данным автора (options.notes)."""
    items = _content_items(document)
    objects = _find_objects(items)
    variants = _variants(objects)
    added = 0
    for obj in objects:
        state = _reference_state(items, obj)
        if state == "colon":
            _add_reference_before_colon(items[obj.first - 1], obj)
        elif state == "missing":
            sentence = _reference_sentence(obj, variants[obj.first])
            previous = items[obj.first - 1] if obj.first else None
            gap = Pt(0) if previous is None or _is_plain_text(previous) or _is_heading(previous) else BLANK_LINE
            if sentence:
                _element(items[obj.first]).addprevious(_new_paragraph(document, sentence, gap)._p)
                added += 1
        if not _has_text_after(items, obj):
            note = options.notes.get(_object_key(document, obj), "").strip()
            if note:
                _element(items[obj.last]).addnext(_new_paragraph(document, note)._p)
                added += 1
    return added


def _center(paragraph, bold_title: bool = False) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.left_indent = Cm(0)
    paragraph.paragraph_format.keep_with_next = True
    if bold_title:
        for run in paragraph.runs:
            run.bold = True


def fix_appendix_headings(document) -> int:
    """Приложение: «ПРИЛОЖЕНИЕ А», ниже по центру статус в скобках, ещё ниже заголовок приложения по центру полужирным."""
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    fixed = 0
    for index, paragraph in enumerate(paragraphs):
        if not APPENDIX_RE.match(paragraph.text.strip()):
            continue
        _center(paragraph)
        fixed += 1
        following = paragraphs[index + 1 : index + 3]
        if following and APPENDIX_STATUS_RE.match(following[0].text.strip()):
            _center(following[0])
            following = following[1:]
        if following and len(following[0].text) <= APPENDIX_TITLE_MAX_CHARS and not _is_caption(following[0]):
            _center(following[0], bold_title=True)
    return fixed


def _find_sources_heading(paragraphs: list[Paragraph]) -> int | None:
    return next((i for i, p in enumerate(paragraphs) if _is_heading(p) and p.text.strip().lower() == SOURCES_TITLE), None)


def _entry_paragraphs(paragraphs: list[Paragraph], heading_index: int) -> list[Paragraph]:
    entries = []
    for paragraph in paragraphs[heading_index + 1 :]:
        if _is_heading(paragraph):
            break
        if paragraph.text.strip():
            entries.append(paragraph)
    if entries and entries[0].text.strip().endswith(":"):
        entries[0]._p.getparent().remove(entries[0]._p)
        entries = entries[1:]
    return entries


def _insert_copy_after(paragraph: Paragraph) -> Paragraph:
    clone = copy.deepcopy(paragraph._p)
    paragraph._p.addnext(clone)
    return Paragraph(clone, paragraph._parent)


def _style_entry(paragraph: Paragraph) -> None:
    for run in paragraph.runs:
        run.italic = False
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(standard.PARAGRAPH_INDENT_CM)
    paragraph.paragraph_format.left_indent = Cm(0)
    paragraph.paragraph_format.right_indent = Cm(0)


ZERO_WIDTH_SPACE = "​"
URL_RE = re.compile(r"https?://[^\s​]+")


def breakable_urls(text: str) -> str:
    """В адресах после «/», «-», «_», «?», «&», «=» в пути допускается перенос строки: так выравнивание по ширине не растягивает строку."""

    def split(match: re.Match[str]) -> str:
        host, path = re.match(r"(https?://[^/]+)(.*)", match.group(), re.DOTALL).groups()
        return host + re.sub(r"([/\-_?&=])(?=[^/])", lambda m: m.group(1) + ZERO_WIDTH_SPACE, path)

    return URL_RE.sub(split, text.replace(ZERO_WIDTH_SPACE, ""))


def _rebuild_entries(entries: list[Paragraph], options: FixOptions) -> dict[int, tuple[int, ...]]:
    """Приводит записи к виду ВАК, нумерует по порядку и возвращает соответствие старых номеров в тексте новым."""
    unwanted = [re.compile(pattern, re.IGNORECASE) for pattern in options.drop_sources]
    mapping: dict[int, tuple[int, ...]] = {number: () for number in options.drop_citations}
    next_number = 1
    for position, paragraph in enumerate(entries, 1):
        text_number = position + options.citation_offset
        if any(pattern.search(paragraph.text) for pattern in unwanted):
            paragraph._p.getparent().remove(paragraph._p)
            mapping[text_number] = ()
            continue
        produced = vak_entries(LEADING_MARK_RE.sub("", paragraph.text).strip())
        current = paragraph
        for index, text in enumerate(produced):
            if index:
                current = _insert_copy_after(current)
            _set_text(current, f"{next_number + index} {breakable_urls(text)}")
            _style_entry(current)
        mapping[text_number] = tuple(range(next_number, next_number + len(produced)))
        next_number += len(produced)
    return mapping


def _all_paragraphs(document) -> list[Paragraph]:
    return [Paragraph(element, document) for element in document.element.body.iter(qn("w:p"))]


def _trim_space_before_punctuation(paragraph: Paragraph) -> None:
    runs = [run for run in paragraph.runs if run.text]
    for current, following in zip(runs, runs[1:], strict=False):
        if current.text.endswith(" ") and following.text[:1] in ",.;:)":
            current.text = current.text.rstrip(" ")


def renumber_citations(paragraphs: list[Paragraph], mapping: dict[int, tuple[int, ...]]) -> int:
    """Перенумеровывает ссылки на источники в переданных абзацах, сохраняя оформление фрагментов."""
    changed = 0
    for paragraph in paragraphs:
        if "[" not in paragraph.text:
            continue
        before = paragraph.text
        for run in paragraph.runs:
            if "[" in run.text:
                run.text = renumber(run.text, mapping)
        if paragraph.text == before and renumber(before, mapping) != before:
            _set_text(paragraph, renumber(before, mapping))
        if paragraph.text != before:
            _trim_space_before_punctuation(paragraph)
            changed += 1
    return changed


def _highest_citation(paragraphs: list[Paragraph]) -> int:
    numbers = [n for paragraph in paragraphs for match in CITATION_RE.finditer(paragraph.text) for n in parse_numbers(match.group("numbers"))]
    return max(numbers, default=0)


def _options_for_list(options: FixOptions, entry_count: int, scope: list[Paragraph], several_lists: bool) -> FixOptions:
    """В документе с несколькими списками сдвиг номеров в тексте определяется по каждому списку отдельно."""
    if not several_lists or options.citation_offset or options.drop_citations:
        return options
    offset = max(0, _highest_citation(scope) - entry_count)
    return replace(options, citation_offset=offset, drop_citations=tuple(range(1, offset + 1)))


def fix_sources_list(document, options: FixOptions = FixOptions()) -> int:
    """Источники: без вводной фразы, тире и курсива; сайты по ВАК РБ; ненужные записи удалены; ссылки в тексте каждого раздела согласованы с номерами его списка."""
    paragraphs = _all_paragraphs(document)
    headings = [i for i, p in enumerate(paragraphs) if _is_heading(p) and p.text.strip().lower() == SOURCES_TITLE]
    total = 0
    scope_start = 0
    skip: set = set()
    for heading in headings:
        entries = _entry_paragraphs(paragraphs, heading)
        scope = [p for p in paragraphs[scope_start:heading] if p._p not in skip]
        skip |= {entry._p for entry in entries}
        list_options = _options_for_list(options, len(entries), scope, len(headings) > 1)
        mapping = _rebuild_entries(entries, list_options)
        renumber_citations(scope, mapping)
        total += sum(len(numbers) for numbers in mapping.values())
        scope_start = heading + 1
    return total


LATIN_SPAN_RE = re.compile(r"[A-Za-z][A-Za-z0-9]*(?:(?:://|[._\-/'’+&#]+)[A-Za-z0-9]+| [A-Za-z][A-Za-z0-9]*)*")
LATIN_LETTER_RE = re.compile(r"[A-Za-z]")
TEXT_ONLY_CHILDREN = {qn("w:rPr"), qn("w:t"), qn("w:lastRenderedPageBreak")}


def _is_italic(paragraph: Paragraph, run: Run) -> bool:
    if run.italic is not None:
        return run.italic
    if run.style is not None and run.style.font.italic is not None:
        return run.style.font.italic
    style = paragraph.style
    while style is not None:
        if style.font.italic is not None:
            return style.font.italic
        style = style.base_style
    return False


def _latin_segments(text: str) -> list[tuple[str, bool]]:
    """Делит текст на латинские фрагменты (курсив) и всё остальное: кириллицу, цифры, знаки препинания, пробелы (прямой шрифт)."""
    segments, position = [], 0
    for match in LATIN_SPAN_RE.finditer(text):
        if match.start() > position:
            segments.append((text[position : match.start()], False))
        segments.append((match.group(), True))
        position = match.end()
    if position < len(text):
        segments.append((text[position:], False))
    return segments


def _split_run_by_script(paragraph: Paragraph, run: Run) -> int:
    segments = _latin_segments(run.text)
    if all(is_latin for _, is_latin in segments):
        return 0
    if not LATIN_LETTER_RE.search(run.text):
        run.italic = False
        return 1
    if {child.tag for child in run._r} - TEXT_ONLY_CHILDREN:
        return 0
    for text, is_latin in segments:
        clone = copy.deepcopy(run._r)
        for marker in clone.findall(qn("w:lastRenderedPageBreak")):
            clone.remove(marker)
        run._r.addprevious(clone)
        piece = Run(clone, paragraph)
        piece.text = text
        piece.italic = is_latin
    run._r.getparent().remove(run._r)
    return 1


def _document_paragraphs(document) -> list[Paragraph]:
    parents = [document.element.body] + [part._element for section in document.sections for part in (section.header, section.footer)]
    return [Paragraph(element, document) for parent in parents for element in parent.iter(qn("w:p"))]


def fix_italics(document) -> int:
    """Курсивом набирается только латиница: кириллица, цифры, знаки препинания и пробелы остаются прямыми, в том числе рядом с латиницей."""
    fixed = 0
    for paragraph in _document_paragraphs(document):
        for element in paragraph._p.xpath(".//w:r"):
            run = Run(element, paragraph)
            if run.text and _is_italic(paragraph, run):
                fixed += _split_run_by_script(paragraph, run)
    return fixed


LIST_MARK_RE = re.compile(r"^[–—\-•]\s+")


def fix_numbered_lists(document) -> int:
    """Перечисления с тире или маркером заменяются нумерованным списком с цифрами: «1) …»; нумерация идёт заново в каждом списке."""
    fixed = number = 0
    started = False
    for paragraph in document.paragraphs:
        text = paragraph.text
        name = paragraph.style.name
        started = started or (name.startswith("Heading") and bool(text.strip()))
        if not text.strip():
            continue
        if not started or name.startswith(SKIPPED_STYLE_PREFIXES) or not LIST_MARK_RE.match(text):
            number = 0
            continue
        number += 1
        run = next((run for run in paragraph.runs if run.text), None)
        if run is not None and LIST_MARK_RE.match(run.text):
            run.text = LIST_MARK_RE.sub(f"{number}) ", run.text, count=1)
            fixed += 1
    return fixed


def count_non_latin_italic(document) -> int:
    """Фрагменты курсивом, в которых есть что-то кроме латиницы: кириллица, цифры, знаки препинания."""
    count = 0
    for paragraph in _document_paragraphs(document):
        for element in paragraph._p.xpath(".//w:r"):
            run = Run(element, paragraph)
            if run.text.strip() and _is_italic(paragraph, run) and not all(is_latin for _, is_latin in _latin_segments(run.text.strip())):
                count += 1
    return count


SKIPPED_STYLE_PREFIXES = ("Heading", "toc", "TOC", "Title")


def _is_centered(paragraph: Paragraph) -> bool:
    value = paragraph.alignment
    style = paragraph.style
    while value is None and style is not None:
        value = style.paragraph_format.alignment
        style = style.base_style
    return value == WD_ALIGN_PARAGRAPH.CENTER


def fix_body_indents(document) -> int:
    """Основной текст, списки и записи источников: абзацный отступ 1,25 см, слева и справа без отступов, по ширине.

    Титульный лист, заголовки, содержание, подписи, рисунки и текст по центру не меняются.
    """
    fixed = 0
    started = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        name = paragraph.style.name
        started = started or (name.startswith("Heading") and bool(text))
        if not started or not text or name.startswith(SKIPPED_STYLE_PREFIXES) or _has_picture(paragraph) or _is_caption(paragraph):
            continue
        if CONTINUATION_RE.match(text) or _is_centered(paragraph):
            continue
        fmt = paragraph.paragraph_format
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fmt.left_indent = Cm(0)
        fmt.right_indent = Cm(0)
        fmt.first_line_indent = Cm(standard.PARAGRAPH_INDENT_CM)
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.line_spacing = Pt(standard.LINE_SPACING_PT)
        fixed += 1
    return fixed


def fix_all_objects(document, options: FixOptions = FixOptions()) -> int:
    return (
        fix_title_date(document)
        + remove_filler_paragraphs(document)
        + fix_section_headings(document)
        + fix_subsection_spacing(document)
        + fix_text_around_objects(document, options)
        + fix_objects(document)
        + fix_appendix_headings(document)
        + fix_sources_list(document, options)
        + fix_numbered_lists(document)
        + fix_italics(document)
        + fix_body_indents(document)
    )
