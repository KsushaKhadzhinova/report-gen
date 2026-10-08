from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt
from docx.text.paragraph import Paragraph

from reportgen.domain import enterprise_standard as standard
from reportgen.domain.citations import CITATION_RE, parse_numbers
from reportgen.domain.fix_options import FixOptions
from reportgen.domain.lint_rules import Issue
from reportgen.infrastructure.docx_objects import (
    count_objects_without_reference,
    count_objects_without_text,
    count_wrong_italics,
    figure_image,
    fix_all_objects,
    objects_needing_notes,
    sources_entries,
)

MIN_BODY_PARAGRAPH_CHARS = 80
MARGIN_TOLERANCE = Mm(1)
TABLE_CAPTION_RE = re.compile(r"^Таблица\s+[\dА-Я]+(\.\d+)?\s*[–-]")
NUMBERED_TITLE_RE = re.compile(r"^\d+(\.\d+)*\s")
APPENDIX_HEADING_RE = re.compile(r"^ПРИЛОЖЕНИЕ\s+[А-ЯA-Z]\b")
CAPTION_RE = re.compile(r"^(Рисунок|Таблица|Листинг)\s+[\dА-Я]+(\.\d+)?\s*[–-]")
INDENT_TOLERANCE = Cm(0.1)


def _inherited(paragraph, attribute: str):
    """Действующее значение свойства абзаца: собственное, а если не задано, то из стиля и его родителей."""
    value = getattr(paragraph.paragraph_format if attribute != "alignment" else paragraph, attribute)
    style = paragraph.style
    while value is None and style is not None:
        value = getattr(style.paragraph_format, attribute)
        style = style.base_style
    return value


def _alignment(paragraph):
    return _inherited(paragraph, "alignment") or WD_ALIGN_PARAGRAPH.LEFT


def _is_caption(paragraph) -> bool:
    return CAPTION_RE.match(paragraph.text.strip()) is not None


def _is_listing(paragraph) -> bool:
    return paragraph.style.name == standard.LISTING_STYLE


def _is_body(paragraph) -> bool:
    long_enough = len(paragraph.text) >= MIN_BODY_PARAGRAPH_CHARS
    centered = _alignment(paragraph) == WD_ALIGN_PARAGRAPH.CENTER
    name = paragraph.style.name.lower()
    contents_entry = name.startswith(("toc", "оглавление"))
    return (
        long_enough
        and not centered
        and not name.startswith("heading")
        and not contents_entry
        and not _is_caption(paragraph)
        and not _is_listing(paragraph)
    )


def _has_foreign_font(paragraph) -> bool:
    return any(
        run.font.name not in (None, standard.FONT) or run.font.size not in (None, Pt(standard.FONT_SIZE_PT)) for run in paragraph.runs
    )


def _has_wrong_indent(paragraph) -> bool:
    indent = _inherited(paragraph, "first_line_indent")
    return indent is None or abs(indent - Cm(standard.PARAGRAPH_INDENT_CM)) > INDENT_TOLERANCE


def _has_side_indent(paragraph) -> bool:
    return any(abs(_inherited(paragraph, side) or 0) > INDENT_TOLERANCE for side in ("left_indent", "right_indent"))


def _is_heading(paragraph) -> bool:
    return paragraph.style.name.startswith("Heading") and bool(paragraph.text.strip())


def _heading_alignment_is_wrong(paragraph) -> bool:
    appendix = APPENDIX_HEADING_RE.match(paragraph.text.strip()) is not None
    centered_expected = appendix or (paragraph.style.name == "Heading 1" and NUMBERED_TITLE_RE.match(paragraph.text.strip()) is None)
    expected = WD_ALIGN_PARAGRAPH.CENTER if centered_expected else WD_ALIGN_PARAGRAPH.LEFT
    return _alignment(paragraph) != expected


def _is_table_caption(paragraph) -> bool:
    return TABLE_CAPTION_RE.match(paragraph.text.strip()) is not None


SOURCES_TITLE = "список использованных источников"


def _list_entries(paragraphs: list[Paragraph], heading_index: int) -> list[Paragraph]:
    block = []
    for paragraph in paragraphs[heading_index + 1 :]:
        if _is_heading(paragraph):
            break
        if paragraph.text.strip():
            block.append(paragraph)
    return block


def _citation_issues(document, name: str) -> list[Issue]:
    """Ссылки в тексте на номера, которых нет в соответствующем списке источников (у каждого раздела свой список)."""
    paragraphs = [Paragraph(element, document) for element in document.element.body.iter(qn("w:p"))]
    headings = [i for i, p in enumerate(paragraphs) if _is_heading(p) and p.text.strip().lower() == SOURCES_TITLE]
    issues: list[Issue] = []
    scope_start = 0
    listed: set[int] = set()
    for heading in headings:
        block = _list_entries(paragraphs, heading)
        entries = block[1:] if block and block[0].text.strip().endswith(":") else block
        listed |= {id(p._p) for p in block}
        cited = {
            number
            for paragraph in paragraphs[scope_start:heading]
            if id(paragraph._p) not in listed
            for match in CITATION_RE.finditer(paragraph.text)
            for number in parse_numbers(match.group("numbers"))
        }
        missing = sorted(number for number in cited if number > len(entries))
        if missing:
            issues.append(Issue("error", name, f"ссылки на источники, которых нет в списке ({len(entries)} записей): {missing}"))
        scope_start = heading + 1
    return issues


def _margin_issues(document, name: str) -> list[Issue]:
    section = document.sections[0]
    expected = {
        "левое поле": (section.left_margin, standard.MARGIN_LEFT_MM),
        "правое поле": (section.right_margin, standard.MARGIN_RIGHT_MM),
        "верхнее поле": (section.top_margin, standard.MARGIN_TOP_MM),
        "нижнее поле": (section.bottom_margin, standard.MARGIN_BOTTOM_MM),
    }
    return [
        Issue("error", name, f"{label}: ожидается {wanted_mm} мм")
        for label, (actual, wanted_mm) in expected.items()
        if actual is None or abs(actual - Mm(wanted_mm)) > MARGIN_TOLERANCE
    ]


class _DocxText:
    def __init__(self, path: Path) -> None:
        self.path = path
        self.document = Document(str(path))

    def paragraphs(self) -> list[str]:
        return [p.text for p in self.document.paragraphs if p.text.strip()]

    def replace(self, old: str, new: str) -> None:
        for paragraph in self.document.paragraphs:
            if paragraph.text == old:
                for run in paragraph.runs[1:]:
                    run._r.getparent().remove(run._r)
                if paragraph.runs:
                    paragraph.runs[0].text = new
                else:
                    paragraph.add_run(new)

    def save(self) -> None:
        self.document.save(str(self.path))


class DocxFormatService:
    def audit(self, path: Path) -> list[Issue]:
        document = Document(str(path))
        issues = _margin_issues(document, path.name) + _citation_issues(document, path.name)
        if document.styles["Normal"].font.name != standard.FONT:
            issues.append(Issue("error", path.name, f"Шрифт Normal не {standard.FONT}"))
        missing_text = count_objects_without_text(document)
        if missing_text:
            issues.append(Issue("warning", path.name, f"без текста после таблицы или рисунка: {missing_text}"))
        missing_reference = count_objects_without_reference(document)
        if missing_reference:
            issues.append(Issue("warning", path.name, f"без текста со ссылкой перед таблицей или рисунком: {missing_reference}"))
        entries = sources_entries(document)
        listed = {id(entry._p) for entry in entries}
        body = [p for p in document.paragraphs if _is_body(p) and id(p._p) not in listed]
        counts = {
            "абзацев с другим шрифтом или размером": sum(_has_foreign_font(p) for p in body),
            "абзацев с абзацным отступом не 1,25 см": sum(_has_wrong_indent(p) for p in body),
            "абзацев не по ширине": sum(_alignment(p) != WD_ALIGN_PARAGRAPH.JUSTIFY for p in body),
            "абзацев с отступом слева или справа": sum(_has_side_indent(p) for p in body),
            "фрагментов, где курсив не совпадает с латиницей": count_wrong_italics(document),
        }
        counts["заголовков с неверным выравниванием"] = sum(_heading_alignment_is_wrong(p) for p in document.paragraphs if _is_heading(p))
        counts["подписей таблиц не по левому краю"] = sum(
            _alignment(p) != WD_ALIGN_PARAGRAPH.LEFT for p in document.paragraphs if _is_table_caption(p)
        )
        issues += [Issue("error", path.name, f"{count} {label}") for label, count in counts.items() if count]
        return issues

    def fix(self, source: Path, output: Path, options: FixOptions = FixOptions()) -> int:
        document = Document(str(source))
        self._fix_page(document)
        self._fix_normal_style(document)
        fixed = sum(self._fix_paragraph(p) for p in document.paragraphs)
        fixed += sum(self._fix_heading(p) for p in document.paragraphs if _is_heading(p))
        fixed += sum(self._fix_table_caption(p) for p in document.paragraphs if _is_table_caption(p))
        fixed += fix_all_objects(document, options)
        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output))
        return fixed

    def open_text(self, path: Path) -> _DocxText:
        return _DocxText(path)

    def notes_template(self, path: Path, images_dir: Path | None = None) -> list[dict]:
        """Таблицы и рисунки, после которых нужен абзац: данные для автора, который пишет текст; файлы рисунков кладутся в images_dir."""
        document = Document(str(path))
        entries = objects_needing_notes(document)
        if images_dir is not None:
            images_dir.mkdir(parents=True, exist_ok=True)
            for entry in entries:
                image = figure_image(document, entry["key"]) if entry["kind"] == "figure" else None
                if image:
                    (images_dir / image[0]).write_bytes(image[1])
                    entry["image"] = str(images_dir / image[0])
        return entries

    @staticmethod
    def _fix_heading(paragraph) -> int:
        fmt = paragraph.paragraph_format
        centered = paragraph.style.name == "Heading 1" and NUMBERED_TITLE_RE.match(paragraph.text.strip()) is None
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
        fmt.first_line_indent = None if centered else Cm(standard.PARAGRAPH_INDENT_CM)
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.line_spacing = Pt(standard.LINE_SPACING_PT)
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(standard.LINE_SPACING_PT)
        fmt.keep_with_next = True
        for run in paragraph.runs:
            run.font.name = standard.FONT
            run.font.size = Pt(standard.FONT_SIZE_PT)
            run.font.bold = True
        return 1

    @staticmethod
    def _fix_table_caption(paragraph) -> int:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.keep_with_next = True
        return 1

    @staticmethod
    def _fix_page(document) -> None:
        for section in document.sections:
            section.left_margin = Mm(standard.MARGIN_LEFT_MM)
            section.right_margin = Mm(standard.MARGIN_RIGHT_MM)
            section.top_margin = Mm(standard.MARGIN_TOP_MM)
            section.bottom_margin = Mm(standard.MARGIN_BOTTOM_MM)

    @staticmethod
    def _fix_normal_style(document) -> None:
        normal = document.styles["Normal"]
        normal.font.name = standard.FONT
        normal.font.size = Pt(standard.FONT_SIZE_PT)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), standard.FONT)

    @staticmethod
    def _fix_paragraph(paragraph) -> int:
        if paragraph.style.name.startswith("Heading") or not paragraph.text.strip() or _is_listing(paragraph):
            return 0
        for run in paragraph.runs:
            run.font.name = standard.FONT
            run.font.size = Pt(standard.FONT_SIZE_PT)
        if not _is_body(paragraph):
            return 0
        fmt = paragraph.paragraph_format
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fmt.first_line_indent = Cm(standard.PARAGRAPH_INDENT_CM)
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.line_spacing = Pt(standard.LINE_SPACING_PT)
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        return 1
