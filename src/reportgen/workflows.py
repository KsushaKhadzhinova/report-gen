from __future__ import annotations

import re
from pathlib import Path

import yaml
from docx import Document

from reportgen import overlap, style
from reportgen.lint import audit_docx, fix_docx
from reportgen.llm import LLM
from reportgen.readers import read_paragraphs

NUMBERED_HEADING_RE = re.compile(r"^(\d+(?:\.\d+){0,2})\.?\s+(\S.{2,120})$")
PLAIN_HEADINGS = ("ВВЕДЕНИЕ", "ЗАКЛЮЧЕНИЕ", "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", "ЦЕЛЬ РАБОТЫ", "ВЫВОДЫ")
MIN_WORDS, MAX_WORDS = 80, 900


def read_task(path: Path, limit: int = 5000) -> str:
    return "\n".join(read_paragraphs(path))[:limit]


def _slug(index: int, title: str) -> str:
    latin = re.sub(r"[^a-z0-9]+", "_", title.lower())[:20].strip("_")
    return f"{index:02d}_{latin or 'section'}"


def structure_from_sample(sample: Path, name: str = "Структура по образцу") -> dict:
    """Извлекает из образца только заголовки и приблизительные объёмы, текст образца не используется."""
    entries: list[dict] = []

    if sample.suffix.lower() == ".docx":
        for paragraph in Document(str(sample)).paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            if paragraph.style.name.startswith("Heading"):
                level = int(paragraph.style.name.split()[-1]) if paragraph.style.name.split()[-1].isdigit() else 1
                entries.append({"title": text, "level": level, "words": 0})
            elif entries:
                entries[-1]["words"] += len(text.split())
    else:
        for line in read_paragraphs(sample):
            match = NUMBERED_HEADING_RE.match(line)
            if match and len(line) < 140:
                entries.append({"title": match.group(2), "level": match.group(1).count(".") + 1, "words": 0})
            elif line.upper() in PLAIN_HEADINGS:
                entries.append({"title": line.upper(), "level": 1, "words": 0})
            elif entries:
                entries[-1]["words"] += len(line.split())

    sections = []
    for index, entry in enumerate(entries, 1):
        title = re.sub(r"^\d+(\.\d+)*\.?\s+", "", entry["title"])
        has_body = entry["words"] > 0
        sections.append(
            {
                "id": _slug(index, title),
                "title": title,
                "level": entry["level"],
                "numbered": title.upper() not in PLAIN_HEADINGS,
                "words": max(MIN_WORDS, min(MAX_WORDS, entry["words"])) if has_body else 0,
                "guide": f"Раскрой тему «{title}» применительно к проекту.",
            }
        )
    if not sections:
        raise ValueError("В образце не найдено заголовков")
    return {"name": name, "sections": sections}


def save_structure(structure: dict, project: Path) -> Path:
    target = project / "structure.yaml"
    target.write_text(yaml.safe_dump(structure, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return target


def fix_report(
    llm: LLM | None,
    source: Path,
    output: Path,
    reference: Path | None = None,
    remarks: Path | None = None,
) -> dict:
    """Исправляет оформление готовой записки, переписывает совпадающие абзацы и разбирает замечания."""
    before = audit_docx(source)
    report: dict = {"issues_before": [str(i) for i in before]}
    report.update(fix_docx(source, output))

    if reference and llm:
        document = Document(str(output))
        flagged = {m.text for m in overlap.scan([p.text for p in document.paragraphs if len(p.text) > 120], reference)}
        rewritten = 0
        for paragraph in document.paragraphs:
            if paragraph.text in flagged:
                new_text = llm.chat(
                    "rewriter",
                    style.style_prompt(),
                    "Перепиши абзац другими словами, сохранив смысл и факты. Верни только абзац.\n\n" + paragraph.text,
                    temperature=0.9,
                )
                for run in paragraph.runs[1:]:
                    run._r.getparent().remove(run._r)
                if paragraph.runs:
                    paragraph.runs[0].text = new_text
                else:
                    paragraph.add_run(new_text)
                rewritten += 1
        document.save(str(output))
        report["paragraphs_rewritten"] = rewritten

    if remarks and llm:
        text = "\n".join(p.text for p in Document(str(output)).paragraphs if p.text.strip())[:9000]
        review = llm.chat(
            "reviewer",
            "Ты редактор пояснительных записок по СТП БГУИР. Отвечай по-русски, конкретно, списком правок.",
            f"Замечания руководителя:\n{remarks.read_text(encoding='utf-8')}\n\nТекст записки:\n{text}\n\n"
            "Для каждого замечания укажи, в каком месте текста оно применимо и как именно исправить.",
            temperature=0.3,
        )
        review_file = output.with_suffix(".review.md")
        review_file.write_text(review, encoding="utf-8")
        report["review"] = str(review_file)

    report["issues_after"] = [str(i) for i in audit_docx(output)]
    return report
