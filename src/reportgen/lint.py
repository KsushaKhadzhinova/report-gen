from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

from reportgen import stp
from reportgen.document import Block

FIRST_PERSON_RE = re.compile(r"(?<![а-яё])(я|мы|наш\w*|мой|моя)(?![а-яё])", re.IGNORECASE)
PLACEHOLDER_RE = re.compile(r"TODO|\[нет файла|\?\?\?|lorem ipsum", re.IGNORECASE)


@dataclass
class Issue:
    level: str
    where: str
    message: str

    def __str__(self) -> str:
        return f"[{self.level}] {self.where}: {self.message}"


def check_blocks(blocks: list[Block]) -> list[Issue]:
    issues: list[Issue] = []
    text = " ".join(b.text for b in blocks if b.kind == "paragraph").lower()
    previous_level = 0

    for block in blocks:
        where = block.source or "документ"
        if block.kind == "heading":
            if block.level > previous_level + 1 and previous_level:
                issues.append(Issue("error", where, f"Пропущен уровень заголовка перед «{block.text}»"))
            previous_level = block.level
            if block.level == 1 and block.numbered and not block.appendix and block.text.upper() in stp.UNNUMBERED_HEADINGS:
                issues.append(Issue("error", where, f"Заголовок «{block.text}» не нумеруется: добавьте {{-}}"))
        elif block.kind == "paragraph":
            if FIRST_PERSON_RE.search(block.text):
                issues.append(Issue("warning", where, f"Личные местоимения: «{block.text[:60]}…»"))
            if "  " in block.text or "---" in block.text:
                issues.append(Issue("warning", where, f"Двойной пробел или «---»: «{block.text[:60]}…»"))
            if PLACEHOLDER_RE.search(block.text):
                issues.append(Issue("error", where, f"Заглушка в тексте: «{block.text[:60]}…»"))
        elif block.kind == "figure":
            if not block.caption:
                issues.append(Issue("error", where, f"Рисунок {block.number} без подписи"))
            if not re.search(rf"рисунк\w*\s+{re.escape(block.number)}", text):
                issues.append(Issue("warning", where, f"Нет ссылки на рисунок {block.number}"))
        elif block.kind == "table":
            if not block.caption:
                issues.append(Issue("error", where, "Таблица без заголовка"))
            elif not re.search(rf"таблиц\w*\s+{re.escape(block.number)}", text):
                issues.append(Issue("warning", where, f"Нет ссылки на таблицу {block.number}"))
            if len(block.rows) < 2:
                issues.append(Issue("warning", where, f"Таблица {block.number} без строк данных"))
    return issues


def audit_docx(path: Path) -> list[Issue]:
    doc = Document(str(path))
    issues: list[Issue] = []
    section = doc.sections[0]
    expected = {
        "левое поле": (section.left_margin, Mm(stp.MARGIN_LEFT_MM)),
        "правое поле": (section.right_margin, Mm(stp.MARGIN_RIGHT_MM)),
        "верхнее поле": (section.top_margin, Mm(stp.MARGIN_TOP_MM)),
        "нижнее поле": (section.bottom_margin, Mm(stp.MARGIN_BOTTOM_MM)),
    }
    for name, (actual, wanted) in expected.items():
        if actual is None or abs(actual - wanted) > Mm(1):
            issues.append(Issue("error", path.name, f"{name}: ожидается {wanted.mm:.0f} мм"))

    normal = doc.styles["Normal"]
    if normal.font.name != stp.FONT:
        issues.append(Issue("error", path.name, f"Шрифт Normal: {normal.font.name}, нужен {stp.FONT}"))

    off_font = off_indent = off_align = 0
    for paragraph in doc.paragraphs:
        if len(paragraph.text) < 80 or paragraph.style.name.startswith("Heading"):
            continue
        if any(r.font.name not in (None, stp.FONT) or (r.font.size not in (None, Pt(stp.FONT_SIZE_PT))) for r in paragraph.runs):
            off_font += 1
        indent = paragraph.paragraph_format.first_line_indent
        if indent is None or abs(indent - Cm(stp.PARAGRAPH_INDENT_CM)) > Cm(0.1):
            off_indent += 1
        if paragraph.alignment not in (None, WD_ALIGN_PARAGRAPH.JUSTIFY):
            off_align += 1
    for count, message in (
        (off_font, "абзацев с другим шрифтом или размером"),
        (off_indent, "абзацев с абзацным отступом не 1,25 см"),
        (off_align, "абзацев не по ширине"),
    ):
        if count:
            issues.append(Issue("error", path.name, f"{count} {message}"))
    return issues


def fix_docx(path: Path, output: Path) -> dict[str, int]:
    """Приводит основной текст и поля документа к требованиям СТП."""
    doc = Document(str(path))
    for section in doc.sections:
        section.left_margin = Mm(stp.MARGIN_LEFT_MM)
        section.right_margin = Mm(stp.MARGIN_RIGHT_MM)
        section.top_margin = Mm(stp.MARGIN_TOP_MM)
        section.bottom_margin = Mm(stp.MARGIN_BOTTOM_MM)

    normal = doc.styles["Normal"]
    normal.font.name = stp.FONT
    normal.font.size = Pt(stp.FONT_SIZE_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), stp.FONT)

    fixed = 0
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading") or not paragraph.text.strip():
            continue
        for run in paragraph.runs:
            run.font.name = stp.FONT
            run.font.size = Pt(stp.FONT_SIZE_PT)
        if len(paragraph.text) >= 80 and paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER:
            fmt = paragraph.paragraph_format
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            fmt.first_line_indent = Cm(stp.PARAGRAPH_INDENT_CM)
            fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            fmt.line_spacing = Pt(stp.LINE_SPACING_PT)
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(0)
            fixed += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return {"paragraphs_fixed": fixed}
