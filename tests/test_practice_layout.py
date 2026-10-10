from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from reportgen.domain.fix_options import FixOptions
from reportgen.infrastructure.docx_format import DocxFormatService
from reportgen.infrastructure.docx_practice import (
    city_at_page_bottom,
    dash_lists,
    frame_screenshots,
    insert_contents,
    lighten_dark_screenshots,
    listings_to_appendices,
    rename_conclusion,
    start_sections_on_new_page,
)

BODY = "Это обычный абзац основного текста отчёта, достаточно длинный, чтобы считаться текстом раздела."


def build(tmp_path: Path) -> Path:
    document = Document()
    document.add_paragraph("Минск 2026")
    document.add_heading("1 Цель работы", level=1)
    document.add_paragraph(BODY)
    document.add_heading("2 Задание", level=1)
    document.add_paragraph("создать приложение", style="List Bullet")
    document.add_paragraph("добавить обработку ошибок.", style="List Bullet")
    document.add_heading("3 Ход работы", level=1)
    document.add_heading("3.1 Настройка проекта", level=2)
    document.add_paragraph("Листинг 1 – Настройка приложения (app.js)")
    document.add_paragraph("const app = express();")
    document.add_paragraph("module.exports = app;")
    document.add_heading("3.2 Шаблоны", level=2)
    document.add_paragraph("Листинг 2 – Шаблон главной страницы")
    document.add_paragraph("<h1>Список</h1>")
    document.add_heading("5 Выводы", level=1)
    document.add_paragraph(BODY)
    path = tmp_path / "practice.docx"
    document.save(path)
    return path


def texts(document):
    return [p.text for p in document.paragraphs]


def test_every_first_level_heading_starts_a_new_page(tmp_path: Path):
    document = Document(build(tmp_path))
    assert start_sections_on_new_page(document) == 4
    headings = [p for p in document.paragraphs if p.style.name == "Heading 1"]
    assert all(p.paragraph_format.page_break_before for p in headings)


def test_conclusion_is_unnumbered_and_centered(tmp_path: Path):
    document = Document(build(tmp_path))
    assert rename_conclusion(document) == 1
    heading = next(p for p in document.paragraphs if p.text == "ЗАКЛЮЧЕНИЕ")
    assert heading.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_bullets_become_dash_items_with_semicolons_and_a_final_period(tmp_path: Path):
    document = Document(build(tmp_path))
    assert dash_lists(document) == 2
    items = [t for t in texts(document) if t.startswith("– ")]
    assert items == ["– создать приложение;", "– добавить обработку ошибок."]
    first = next(p for p in document.paragraphs if p.text.startswith("– создать"))
    assert first.style.name == "Normal"
    assert first.paragraph_format.left_indent.cm == 0


def test_contents_are_inserted_once_before_the_first_section(tmp_path: Path):
    document = Document(build(tmp_path))
    assert insert_contents(document) == 1
    assert insert_contents(document) == 0
    order = texts(document)
    assert order.index("СОДЕРЖАНИЕ") < order.index("1 Цель работы")
    instructions = [t.text for t in document.element.body.iter(qn("w:instrText"))]
    assert any("TOC" in text for text in instructions)


def test_city_is_anchored_to_the_bottom_margin(tmp_path: Path):
    document = Document(build(tmp_path))
    assert city_at_page_bottom(document) == 1
    frame = document.paragraphs[0]._p.pPr.find(qn("w:framePr"))
    assert frame.get(qn("w:yAlign")) == "bottom"


def test_listings_move_to_appendices_with_italic_code(tmp_path: Path):
    document = Document(build(tmp_path))
    assert listings_to_appendices(document) == 2
    order = texts(document)
    assert "ПРИЛОЖЕНИЕ А" in order and "ПРИЛОЖЕНИЕ Б" in order
    assert "Листинг А.1 – Настройка приложения (app.js)" in order
    assert "Листинг Б.1 – Шаблон главной страницы" in order
    assert order.index("3.1 Настройка проекта") + 1 == order.index("3.2 Шаблоны")
    code = next(p for p in document.paragraphs if p.text == "const app = express();")
    assert all(run.italic for run in code.runs)
    assert order.index("ПРИЛОЖЕНИЕ А") > order.index("5 Выводы")
    name = order[order.index("ПРИЛОЖЕНИЕ А") + 2]
    assert name == "Исходный код: настройка проекта"


def test_screenshots_get_a_frame_and_dark_terminal_becomes_light(tmp_path: Path):
    from PIL import Image

    dark = tmp_path / "dark.png"
    Image.new("RGB", (40, 20), (12, 12, 12)).save(dark)
    document = Document()
    document.add_picture(str(dark))
    assert frame_screenshots(document) == 1
    assert frame_screenshots(document) == 0
    assert lighten_dark_screenshots(document) == 1
    blob = next(r.target_part.blob for r in document.part.rels.values() if "image" in r.reltype)
    import io

    assert Image.open(io.BytesIO(blob)).getpixel((1, 1))[0] > 200


def test_practice_option_runs_the_whole_layout(tmp_path: Path):
    output = tmp_path / "fixed.docx"
    DocxFormatService().fix(build(tmp_path), output, FixOptions(practice=True))
    order = texts(Document(output))
    assert "СОДЕРЖАНИЕ" in order and "ЗАКЛЮЧЕНИЕ" in order and "ПРИЛОЖЕНИЕ А" in order
    assert any(t.startswith("– ") for t in order)
