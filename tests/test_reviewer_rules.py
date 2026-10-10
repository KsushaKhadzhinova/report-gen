from pathlib import Path

from docx import Document

from reportgen.domain.blocks import Block, Kind
from reportgen.domain.lint_rules import check_blocks
from reportgen.domain.reviewer_rules import NBSP, glue_reference_numbers, guillemets
from reportgen.infrastructure.docx_typography import fix_hyperlink_look, fix_quotes_and_reference_spaces
from reportgen.infrastructure.pdf_audit import page_fill_issues


def messages(blocks):
    return [issue.message for issue in check_blocks(blocks)]


def test_straight_quotes_become_guillemets_but_code_is_left_alone():
    assert guillemets('Режим "тест" включён') == "Режим «тест» включён"
    assert guillemets('app.use("/admin", fakeAuth)') == 'app.use("/admin", fakeAuth)'


def test_reference_word_and_number_are_glued():
    assert glue_reference_numbers("на рисунке 4 и в таблице 5") == f"на рисунке{NBSP}4 и в таблице{NBSP}5"
    assert glue_reference_numbers("Рис. 2") == f"Рис.{NBSP}2"
    assert glue_reference_numbers("таблица без номера") == "таблица без номера"


def test_lint_flags_reviewer_remarks():
    blocks = [
        Block(Kind.HEADING, text="Выводы", level=1, source="a"),
        Block(Kind.PARAGRAPH, text='Система также работает "быстро" (см. рисунок 1).', source="b"),
    ]
    found = " | ".join(messages(blocks))
    assert "Заключение" in found
    assert "Слова-паразиты" in found
    assert "Прямые кавычки" in found
    assert "без «см.» и скобок" in found


def test_goal_of_work_heading_must_be_numbered():
    blocks = [Block(Kind.HEADING, text="Цель работы", level=1, numbered=False, source="a")]
    assert any("нумеруется" in m for m in messages(blocks))


def test_docx_quotes_and_reference_spaces_are_fixed_across_runs(tmp_path: Path):
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run('Режим "тест" на рисунке ')
    paragraph.add_run("4").italic = True
    assert fix_quotes_and_reference_spaces(document) >= 2
    assert paragraph.text == f"Режим «тест» на рисунке{NBSP}4"


def test_hyperlink_runs_lose_color_and_underline():
    document = Document()
    paragraph = document.add_paragraph()
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    link = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = "https://example.com"
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    assert fix_hyperlink_look(document) == 1
    assert run.find(qn("w:rPr")).find(qn("w:color")).get(qn("w:val")) == "auto"
    assert run.find(qn("w:rPr")).find(qn("w:u")).get(qn("w:val")) == "none"


def test_pdf_pages_with_too_few_lines_are_reported(tmp_path: Path):
    import pymupdf

    path = tmp_path / "short.pdf"
    pdf = pymupdf.open()
    for lines in (30, 30, 3, 25, 4):
        page = pdf.new_page()
        for index in range(lines):
            page.insert_text((72, 60 + index * 20), f"line {index}")
    pdf.save(path)
    pdf.close()
    reported = [issue.where for issue in page_fill_issues(path)]
    assert reported == ["short.pdf, страница 3"]
