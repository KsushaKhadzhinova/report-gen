from docx import Document

from reportgen.domain.citations import format_numbers, parse_numbers, renumber
from reportgen.domain.fix_options import FixOptions
from reportgen.domain.sources import split_domain_entry
from reportgen.infrastructure.docx_format import DocxFormatService


def test_numbers_are_parsed_and_formatted_back():
    assert parse_numbers("1–3, 7") == [1, 2, 3, 7]
    assert format_numbers([5, 6, 7, 9]) == "5–7, 9"
    assert format_numbers([5, 6]) == "5, 6"


def test_renumber_maps_ranges_and_removes_dropped_citations():
    mapping = {1: (), 2: (), 3: (), 4: (), 5: (1,), 6: (2,), 7: (3,), 8: (4,), 9: tuple(range(5, 15)), 10: (15,)}
    assert renumber("Работа по методикам [1–3].", mapping) == "Работа по методикам."
    assert renumber("Данные сайтов [4–8] и страницы [9], цены [10].", mapping) == "Данные сайтов [1–4] и страницы [5–14], цены [15]."
    assert renumber("Таблица методики [1, табл. 7] описана.", mapping) == "Таблица методики описана."


def test_citation_with_a_page_tail_keeps_the_tail_for_a_single_source():
    assert renumber("См. [5, с. 12].", {5: (1,)}) == "См. [1, с. 12]."


def test_domain_entry_is_split_into_entries_with_urls():
    entry = "Страницы сайтов конкурентов: plantuml.com, mermaid.js.org [Электронный ресурс]. – Дата доступа: 01.10.2026."
    assert split_domain_entry(entry) == [
        "plantuml.com [Электронный ресурс]. – Режим доступа: https://plantuml.com. – Дата доступа: 01.10.2026.",
        "mermaid.js.org [Электронный ресурс]. – Режим доступа: https://mermaid.js.org. – Дата доступа: 01.10.2026.",
    ]
    assert split_domain_entry("Иванов, И. И. Основы анализа.") is None


def build_report_with_sources(tmp_path):
    document = Document()
    document.add_paragraph("Текст со ссылками на методики [1–3] и на данные [4–6].")
    document.add_heading("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", level=1)
    document.add_paragraph("– Свои отчёты ЛР1–ЛР4. – Минск, 2026.")
    document.add_paragraph("– Similarweb Pro: обзор [Электронный ресурс]. – Режим доступа: https://pro.similarweb.com. – Дата доступа: 30.09.2026.")
    document.add_paragraph("– Страницы сайтов: plantuml.com, mermaid.ai [Электронный ресурс]. – Дата доступа: 01.10.2026.")
    path = tmp_path / "report.docx"
    document.save(str(path))
    return path


def test_sources_are_rebuilt_and_in_text_citations_follow(tmp_path):
    options = FixOptions(drop_sources=("Свои отчёты",), citation_offset=3, drop_citations=(1, 2, 3))
    output = tmp_path / "fixed.docx"
    DocxFormatService().fix(build_report_with_sources(tmp_path), output, options)
    texts = [p.text for p in Document(str(output)).paragraphs]
    body = next(t for t in texts if t.startswith("Текст со ссылками"))
    assert body == "Текст со ссылками на методики и на данные [1–3]."
    entries = [t for t in texts if t[:3] in ("[1]", "[2]", "[3]")]
    assert len(entries) == 3
    assert entries[0] == "[1] Similarweb Pro [Электронный ресурс] : обзор. – Режим доступа: https://pro.similarweb.com. – Дата доступа: 30.09.2026."
    assert entries[1].startswith("[2] plantuml.com") and entries[2].startswith("[3] mermaid.ai")


def test_audit_reports_citations_without_a_source(tmp_path):
    document = Document()
    document.add_paragraph("Данные получены из источников [1] и [5].")
    document.add_heading("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", level=1)
    document.add_paragraph("1 Первый источник.")
    document.add_paragraph("2 Второй источник.")
    path = tmp_path / "report.docx"
    document.save(str(path))
    messages = [issue.message for issue in DocxFormatService().audit(path)]
    assert any("которых нет в списке (2 записей): [5]" in message for message in messages)
