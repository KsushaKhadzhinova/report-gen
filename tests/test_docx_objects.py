from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt

from reportgen.domain.fix_options import FixOptions
from reportgen.domain.sources import to_vak_site_entry
from reportgen.infrastructure.docx_format import DocxFormatService


def build_source(tmp_path: Path) -> Path:
    from PIL import Image

    picture_file = tmp_path / "picture.png"
    Image.new("RGB", (60, 40), "white").save(picture_file)
    document = Document()
    document.add_paragraph("Министерство образования Республики Беларусь")
    document.add_paragraph("Минск, октябрь 2026")
    document.add_heading("ЦЕЛЬ РАБОТЫ", level=1)
    document.add_paragraph("Текст перед таблицей, достаточно длинный, чтобы считаться основным абзацем документа по стандарту предприятия.")
    document.add_paragraph("Таблица 1 – Пример")
    table = document.add_table(rows=3, cols=2)
    for row in table.rows:
        for cell in row.cells:
            cell.text = "Ячейка"
    document.add_paragraph("Текст сразу после таблицы без пустой строки, достаточно длинный, чтобы считаться основным абзацем документа.")
    picture_paragraph = document.add_paragraph()
    picture_paragraph.add_run().add_picture(str(picture_file))
    picture_paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    picture_paragraph.paragraph_format.line_spacing = Pt(18)
    document.add_paragraph("Рисунок 1 – Пример рисунка")
    document.add_heading("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", level=1)
    document.add_paragraph("В качестве источников использованы следующие:")
    entry = document.add_paragraph()
    entry.add_run("– Иванов, И. И. ")
    entry.add_run("Основы анализа").italic = True
    document.add_paragraph("– Петров, П. П. Методика расчёта.")
    source = tmp_path / "source.docx"
    document.save(str(source))
    return source


def fixed(tmp_path: Path) -> Document:
    output = tmp_path / "fixed.docx"
    DocxFormatService().fix(build_source(tmp_path), output)
    return Document(str(output))


def find(document, fragment: str):
    return next(p for p in document.paragraphs if fragment in p.text)


def test_title_has_no_month(tmp_path: Path):
    assert find(fixed(tmp_path), "Минск").text == "Минск 2026"


def test_section_headings_are_capitals_in_text_and_contents(tmp_path: Path):
    document = fixed(tmp_path)
    assert find(document, "ЦЕЛЬ").text == "ЦЕЛЬ РАБОТЫ"


def test_filler_phrase_is_removed(tmp_path: Path):
    from docx import Document as NewDocument

    source = build_source(tmp_path)
    document = NewDocument(str(source))
    document.add_paragraph("В данном подразделе рассматриваются ключевые аспекты темы применительно к продукту.")
    document.save(str(source))
    output = tmp_path / "fixed.docx"
    DocxFormatService().fix(source, output)
    assert not any("В данном подразделе" in p.text for p in NewDocument(str(output)).paragraphs)


def test_site_entry_follows_the_vak_form():
    entry = "Similarweb Pro: обзор доменов [Электронный ресурс]. – Режим доступа: https://pro.similarweb.com. – Дата доступа: 30.09.2026."
    assert to_vak_site_entry(entry) == (
        "Similarweb Pro [Электронный ресурс] : обзор доменов. – Режим доступа: https://pro.similarweb.com. – Дата доступа: 30.09.2026."
    )


def test_resource_mark_goes_after_the_title_proper_before_responsibility():
    entry = "Образование, 2026 : стат. сб. / Нац. стат. ком. [Электронный ресурс]. – Режим доступа: https://www.belstat.gov.by. – Дата доступа: 30.09.2026."
    assert to_vak_site_entry(entry).startswith("Образование, 2026 [Электронный ресурс] : стат. сб. / Нац. стат. ком. – Режим доступа:")


def test_site_entry_without_subtitle_keeps_the_mark_at_the_end_of_the_title():
    entry = "Miro Pricing [Электронный ресурс]. – Режим доступа: https://miro.com/pricing/. – Дата доступа: 30.09.2026."
    assert to_vak_site_entry(entry) == entry


def test_entry_with_url_path_and_slash_keeps_the_address():
    entry = "Sparx Systems: цены [Электронный ресурс]. – Режим доступа: https://sparxsystems.com/products/ea/. – Дата доступа: 02.10.2026."
    assert "Режим доступа: https://sparxsystems.com/products/ea/. – Дата доступа: 02.10.2026." in to_vak_site_entry(entry)


def test_entries_that_are_not_sites_are_left_alone():
    book = "Иванов, И. И. Основы анализа. – Минск : БГУИР, 2025. – 120 с."
    assert to_vak_site_entry(book) == book


def test_unwanted_source_is_dropped_and_the_rest_renumbered(tmp_path: Path):
    output = tmp_path / "fixed.docx"
    DocxFormatService().fix(build_source(tmp_path), output, FixOptions(drop_sources=("Иванов",)))
    texts = [p.text for p in Document(str(output)).paragraphs]
    assert not any("Иванов" in text for text in texts)
    assert next(text for text in texts if "Петров" in text).startswith("[1] Петров")


def test_table_caption_has_a_blank_line_before_and_the_text_after_the_table_too(tmp_path: Path):
    document = fixed(tmp_path)
    assert find(document, "Таблица 1").paragraph_format.space_before == Pt(18)
    assert find(document, "Текст сразу после таблицы").paragraph_format.space_before == Pt(18)


def test_table_uses_standard_font_single_spacing_and_fits_content(tmp_path: Path):
    document = fixed(tmp_path)
    table = document.tables[0]
    paragraph = table.rows[0].cells[0].paragraphs[0]
    assert paragraph.runs[0].font.size == Pt(12)
    assert paragraph.paragraph_format.line_spacing_rule == WD_LINE_SPACING.SINGLE
    assert table.autofit is False
    assert table.rows[0]._tr.trPr.xpath("w:tblHeader")


def test_figure_is_not_clipped_and_is_framed_by_blank_lines(tmp_path: Path):
    document = fixed(tmp_path)
    picture = next(p for p in document.paragraphs if p._p.xpath(".//w:drawing"))
    assert picture.paragraph_format.line_spacing_rule == WD_LINE_SPACING.SINGLE
    assert picture.paragraph_format.space_before == Pt(18)
    assert find(document, "Рисунок 1").paragraph_format.space_after == Pt(18)


def test_sources_have_no_intro_dashes_or_italics_and_are_numbered(tmp_path: Path):
    document = fixed(tmp_path)
    texts = [p.text for p in document.paragraphs]
    assert "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ" in texts
    assert not any("следующие" in text for text in texts)
    first = find(document, "Иванов")
    assert first.text.startswith("[1] Иванов")
    assert all(run.italic is False for run in first.runs)
    assert find(document, "Петров").text.startswith("[2] Петров")
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    assert first.alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_column_widths_keep_short_columns_readable_and_fill_the_text_width():
    from reportgen.infrastructure.docx_objects import TEXT_WIDTH_CM, column_widths_cm

    columns = [
        ["Часть", "A", "B"],
        ["Методика", "Классификация компаний по уровням конкуренции", "Анализ по сайтам"],
        ["Рабочая книга", "Excel/klassifikaciya_urovni_konkurencii_NotaCode.xlsx", "Excel/analiz_konkurentov_Levitt_Kano_NotaCode.xlsx"],
        ["Результат", "Реестр из 22 компаний, оценка по критериям, уровни, карта близости", "Выборка из 8 конкурентов, карта сайтов"],
    ]
    widths = column_widths_cm(columns)
    assert abs(sum(widths) - TEXT_WIDTH_CM) < 0.01
    assert widths[0] >= len("Часть") * 0.22
    assert widths[3] > widths[0] * 2


def build_two_sections(tmp_path: Path) -> Path:
    document = Document()
    for title, text, entries in (
        ("Первый", "Факт [4]. Ещё [5, 6].", ["Метод.", "Сайт А.", "Сайт Б."]),
        ("Второй", "Вывод [1]. Другое [2].", ["Книга.", "Статья."]),
    ):
        document.add_heading(title, level=1)
        document.add_paragraph(text)
        document.add_heading("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", level=1)
        for entry in entries:
            document.add_paragraph(f"– {entry}")
    path = tmp_path / "two.docx"
    document.save(path)
    return path


def test_each_sources_list_gets_its_own_offset_and_the_audit_checks_each_list(tmp_path: Path):
    output = tmp_path / "two_fixed.docx"
    DocxFormatService().fix(build_two_sections(tmp_path), output, FixOptions())
    texts = [p.text for p in Document(output).paragraphs]
    assert "Факт [1]. Ещё [2, 3]." in texts
    assert "Вывод [1]. Другое [2]." in texts
    audit = DocxFormatService().audit(output)
    assert not [issue for issue in audit if "нет в списке" in issue.message]


def test_long_contents_entries_are_not_audited_as_body_paragraphs(tmp_path: Path):
    document = Document()
    document.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("ЛАБОРАТОРНАЯ РАБОТА №5. " + "ОЧЕНЬ ДЛИННОЕ НАЗВАНИЕ " * 5 + "	279", style="toc 1")
    path = tmp_path / "toc.docx"
    document.save(path)
    messages = [issue.message for issue in DocxFormatService().audit(path)]
    assert not any("не по ширине" in m or "абзацным отступом" in m for m in messages)
