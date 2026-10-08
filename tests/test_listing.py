from docx import Document

from reportgen.domain.blocks import Block, Kind
from reportgen.infrastructure.docx_format import DocxFormatService
from reportgen.infrastructure.docx_renderer import DocxRenderer

CODE = "static int Add(int a, int b)\n{\n    return a + b; // сумма\n}\n\npublic class Demo { }"


def _render(tmp_path):
    blocks = [
        Block(Kind.HEADING, text="Выполнение работы", level=1, numbered=True),
        Block(Kind.PARAGRAPH, text="Код метода сложения приведён в листинге, он состоит из нескольких строк и описывает метод целиком."),
        Block(Kind.CODE, text=CODE, caption="Метод Add", number="1.1"),
        Block(Kind.PARAGRAPH, text="После листинга идёт обычный абзац основного текста, который оформляется по общим правилам стандарта."),
    ]
    target = tmp_path / "note.docx"
    DocxRenderer().render(blocks, {"title_page": False}, tmp_path, target)
    return target


def test_listing_uses_listing_style_times_new_roman_12_italic(tmp_path):
    document = Document(str(_render(tmp_path)))
    style = document.styles["Listing"]
    assert style.font.name == "Times New Roman"
    assert style.font.size.pt == 12
    assert style.font.italic is True
    lines = [p for p in document.paragraphs if p.style.name == "Listing"]
    assert [p.text for p in lines][0] == "static int Add(int a, int b)"
    assert len(lines) == len(CODE.splitlines())


def test_fix_and_audit_do_not_touch_listings(tmp_path):
    source = _render(tmp_path)
    fixed = tmp_path / "fixed.docx"
    service = DocxFormatService()
    service.fix(source, fixed)
    messages = [issue.message for issue in service.audit(fixed) if issue.is_error]
    assert not [m for m in messages if "шрифт" in m or "курсив" in m or "ширин" in m or "отступ" in m], messages
    for paragraph in Document(str(fixed)).paragraphs:
        if paragraph.style.name == "Listing":
            assert paragraph.alignment is None or paragraph.alignment == 0
            assert all(run.font.size is None and run.font.italic is None for run in paragraph.runs)
