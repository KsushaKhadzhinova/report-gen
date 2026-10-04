from pathlib import Path

from docx import Document
from docx.shared import Mm

from reportgen.domain.markup import parse
from reportgen.domain.numbering import assign_numbers, resolve_figure_references
from reportgen.infrastructure.docx_format import DocxFormatService
from reportgen.infrastructure.docx_renderer import DocxRenderer
from reportgen.infrastructure.file_project import FileProjectRepository
from reportgen.infrastructure.scaffold import scaffold_project
from reportgen.infrastructure.tex_renderer import TexRenderer, escape

SOURCE = "# АНАЛИЗ\n\n## Предметная область\n\nНа рисунке {fig:er} показана схема.\n\n![Схема](screenshots/er.png)\n"


def sample_blocks():
    result = parse(SOURCE, "s.md")
    assign_numbers(result)
    resolve_figure_references(result)
    return result


def test_docx_has_standard_page_setup_and_passes_audit(tmp_path: Path):
    output = DocxRenderer().render(sample_blocks(), {"title": "Тест"}, tmp_path, tmp_path / "note.docx")
    section = Document(str(output)).sections[0]
    assert abs(section.left_margin - Mm(30)) < Mm(0.5)
    assert abs(section.right_margin - Mm(15)) < Mm(0.5)
    assert not [i for i in DocxFormatService().audit(output) if "поле" in i.message]


def test_formatter_repairs_margins_and_indent(tmp_path: Path):
    source = tmp_path / "bad.docx"
    document = Document()
    document.sections[0].left_margin = Mm(10)
    document.add_paragraph("Длинный абзац основного текста. " * 6)
    document.save(str(source))
    service = DocxFormatService()
    assert service.audit(source)
    assert service.fix(source, tmp_path / "good.docx") == 1
    assert not service.audit(tmp_path / "good.docx")


def test_tex_output_contains_chapter_command_and_escapes(tmp_path: Path):
    tex = TexRenderer().render(sample_blocks(), {"title": "Тест"}, tmp_path, tmp_path / "tex" / "note.tex")
    text = tex.read_text(encoding="utf-8")
    assert r"\chapterheading{1}{АНАЛИЗ}" in text
    assert "нет файла" in text
    assert escape("a_b & 50%") == r"a\_b \& 50\%"


def test_repository_round_trip(tmp_path: Path):
    scaffold_project(tmp_path / "work", title="Тема")
    repository = FileProjectRepository(tmp_path / "work")
    repository.save_section("01", "# ВВЕДЕНИЕ {-}\n\nТекст.\n")
    repository.store_fingerprint("01", "abc")
    assert repository.section_exists("01")
    assert repository.cached_fingerprint("01") == "abc"
    assert repository.read_meta()["title"] == "Тема"
    assert [b.text for b in repository.load_blocks()] == ["ВВЕДЕНИЕ", "Текст."]


def test_builtin_structures_are_valid_yaml():
    import yaml

    from reportgen.domain.structure import Structure
    from reportgen.infrastructure.settings import read_data

    for name in ("coursework", "lab"):
        structure = Structure.from_dict(yaml.safe_load(read_data(f"{name}.yaml")))
        assert structure.sections and all(section.title for section in structure.sections)
