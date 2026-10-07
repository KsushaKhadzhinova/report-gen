import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

from reportgen.domain.fix_options import FixOptions
from reportgen.domain.object_sentences import figure_reference_sentence, table_reference_sentence
from reportgen.infrastructure.docx_format import DocxFormatService
from reportgen.infrastructure.docx_objects import (
    column_widths_cm,
    count_objects_without_reference,
    count_objects_without_text,
    objects_needing_notes,
)

BODY = "Это обычный абзац основного текста отчёта, достаточно длинный, чтобы считаться текстом раздела."
REFERENCED = "Показатели рабочего ядра запросов приведены в таблице 1 и разобраны ниже в этом разделе отчёта."


def build(tmp_path: Path, text_before: str = BODY, with_text_after: bool = False) -> Path:
    document = Document()
    document.add_heading("Раздел", level=1)
    document.add_paragraph(text_before)
    document.add_paragraph("Таблица 1 – Рабочая таблица ядра")
    table = document.add_table(rows=2, cols=2)
    for index, text in enumerate(("А", "Б", "10", "20")):
        table.rows[index // 2].cells[index % 2].text = text
    if with_text_after:
        document.add_paragraph(BODY)
    document.add_heading("Следующий раздел", level=1)
    document.add_paragraph(BODY)
    path = tmp_path / "objects.docx"
    document.save(path)
    return path


def fixed_texts(source: Path, tmp_path: Path, options: FixOptions = FixOptions()) -> list[str]:
    output = tmp_path / "fixed.docx"
    DocxFormatService().fix(source, output, options)
    return [p.text for p in Document(output).paragraphs]


def test_table_without_a_reference_before_it_gets_a_sentence_before_the_caption(tmp_path: Path):
    texts = fixed_texts(build(tmp_path), tmp_path)
    assert texts.index("Данные по теме «рабочая таблица ядра» приведены в таблице 1.") == texts.index("Таблица 1 – Рабочая таблица ядра") - 1


def test_table_with_a_reference_before_it_is_left_alone(tmp_path: Path):
    texts = fixed_texts(build(tmp_path, text_before=REFERENCED), tmp_path)
    assert not [text for text in texts if text.startswith("Данные по теме")]


def test_paragraph_ending_with_a_colon_gets_the_reference_inside(tmp_path: Path):
    texts = fixed_texts(build(tmp_path, text_before=BODY[:-1] + " приведены ниже:"), tmp_path)
    assert any(text.endswith("приведены ниже (таблица 1):") for text in texts)


def test_note_from_the_author_goes_after_the_table_and_without_it_nothing_is_invented(tmp_path: Path):
    source = build(tmp_path)
    key = objects_needing_notes(Document(str(source)))[0]["key"]
    note = "Значение во второй строке вдвое меньше, чем в первой."
    texts = fixed_texts(source, tmp_path, FixOptions(notes={key: note}))
    assert texts[texts.index("СЛЕДУЮЩИЙ РАЗДЕЛ") - 1] == note
    assert note not in fixed_texts(source, tmp_path)


def test_table_followed_by_text_needs_no_note(tmp_path: Path):
    assert objects_needing_notes(Document(str(build(tmp_path, with_text_after=True)))) == []


def test_audit_warns_about_missing_text_and_reference(tmp_path: Path):
    messages = [i.message for i in DocxFormatService().audit(build(tmp_path))]
    assert any("без текста после таблицы или рисунка: 1" in m for m in messages)
    assert any("без текста со ссылкой перед таблицей или рисунком: 1" in m for m in messages)
    document = Document(str(build(tmp_path, text_before=REFERENCED, with_text_after=True)))
    assert count_objects_without_reference(document) == 0
    assert count_objects_without_text(document) == 0


def test_notes_template_has_the_table_data_and_the_reference_sentence(tmp_path: Path):
    entry = DocxFormatService().notes_template(build(tmp_path))[0]
    assert entry["kind"] == "table" and entry["rows"] == [["А", "Б"], ["10", "20"]] and entry["total_rows"] == 2
    assert entry["reference_sentence"].endswith("в таблице 1.") or "таблицу 1" in entry["reference_sentence"]
    json.dumps(entry, ensure_ascii=False)


def test_reference_sentences_keep_abbreviations_and_ignore_other_captions():
    assert table_reference_sentence("Таблица 2.1 – UML-диаграммы") == "Данные по теме «UML-диаграммы» приведены в таблице 2.1."
    assert "рисунке 3" in figure_reference_sentence("Рисунок 3 – Схема данных")
    assert table_reference_sentence("Просто текст") is None


def test_italics_only_for_latin_and_cyrillic_next_to_it_is_upright(tmp_path: Path):
    document = Document()
    document.add_heading("Раздел", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Диаграмма UML-редактор, ERD и Mermaid Chart: схема 40").italic = True
    path = tmp_path / "italics.docx"
    document.save(path)
    output = tmp_path / "italics_fixed.docx"
    DocxFormatService().fix(path, output)
    runs = [run for run in Document(output).paragraphs[1].runs if run.text]
    assert "".join(run.text for run in runs) == "Диаграмма UML-редактор, ERD и Mermaid Chart: схема 40"
    assert {run.text for run in runs if run.italic} == {"UML", "ERD", "Mermaid Chart"}
    assert all(run.italic is False for run in runs if not run.italic)


def test_list_and_quote_paragraphs_have_no_side_indents(tmp_path: Path):
    document = Document()
    document.add_heading("Раздел", level=1)
    for text in ("– Пункт списка", BODY):
        paragraph = document.add_paragraph(text)
        paragraph.paragraph_format.left_indent = Cm(1.25)
        paragraph.paragraph_format.right_indent = Cm(0.85)
    path = tmp_path / "indents.docx"
    document.save(path)
    output = tmp_path / "indents_fixed.docx"
    DocxFormatService().fix(path, output)
    for paragraph in Document(output).paragraphs[1:]:
        fmt = paragraph.paragraph_format
        assert (fmt.left_indent, fmt.right_indent) == (Cm(0), Cm(0))
        assert abs(fmt.first_line_indent - Cm(1.25)) < Cm(0.01) and paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert not [i for i in DocxFormatService().audit(output) if "отступом слева" in i.message]


def test_narrow_columns_keep_their_width_when_the_table_is_too_wide():
    columns = [["№", "16"], ["Дата и сервис", "30.09.2026"]] + [["Файл", "GT_2026-09-30_BY_5y_UML_BPMN_multiTimeline.csv"]] * 3
    widths = column_widths_cm(columns, total=17.0)
    assert widths[0] >= 0.5 and widths[1] >= len("30.09.2026") * 0.25
    assert abs(sum(widths) - 17.0) < 0.01


def test_appendix_status_and_title_are_centered(tmp_path: Path):
    document = Document()
    document.add_paragraph("ПРИЛОЖЕНИЕ Б")
    document.add_paragraph("(обязательное)")
    document.add_paragraph("Контрольная выгрузка Вордстата")
    document.add_paragraph(BODY)
    path = tmp_path / "appendix.docx"
    document.save(path)
    output = tmp_path / "appendix_fixed.docx"
    DocxFormatService().fix(path, output)
    paragraphs = Document(output).paragraphs
    assert [p.alignment for p in paragraphs[:3]] == [WD_ALIGN_PARAGRAPH.CENTER] * 3
    assert all(run.bold for run in paragraphs[2].runs)
    assert paragraphs[3].alignment != WD_ALIGN_PARAGRAPH.CENTER


def test_centered_appendix_heading_of_any_level_passes_the_audit(tmp_path: Path):
    document = Document()
    document.add_heading("ПРИЛОЖЕНИЕ А", level=2)
    path = tmp_path / "appendix_heading.docx"
    document.save(path)
    output = tmp_path / "appendix_heading_fixed.docx"
    DocxFormatService().fix(path, output)
    assert not [i for i in DocxFormatService().audit(output) if "выравниванием" in i.message]


def test_table_split_by_word_counts_text_after_the_last_part(tmp_path: Path):
    document = Document()
    document.add_heading("Раздел", level=1)
    document.add_paragraph(REFERENCED)
    document.add_paragraph("Таблица 1 – Длинная таблица")
    document.add_table(rows=2, cols=1)
    document.add_paragraph("Продолжение таблицы 1")
    document.add_table(rows=2, cols=1)
    document.add_paragraph(BODY)
    assert count_objects_without_text(document) == 0
    assert count_objects_without_reference(document) == 0


def test_dash_lists_become_numbered_lists_that_restart_in_each_list(tmp_path: Path):
    document = Document()
    document.add_heading("Раздел", level=1)
    for text in ("Первый список:", "– Альфа;", "– Бета.", "Второй список:", "– Гамма."):
        document.add_paragraph(text)
    path = tmp_path / "lists.docx"
    document.save(path)
    texts = fixed_texts(path, tmp_path)
    assert texts[1:] == ["Первый список:", "1) Альфа;", "2) Бета.", "Второй список:", "1) Гамма."]


def test_urls_in_sources_can_wrap_without_stretching_the_line():
    from reportgen.infrastructure.docx_objects import ZERO_WIDTH_SPACE, breakable_urls

    text = "Режим доступа: https://www.belstat.gov.by/upload/iblock/6af/a-b.pdf. – Дата"
    wrapped = breakable_urls(text)
    assert wrapped.replace(ZERO_WIDTH_SPACE, "") == text
    assert "/" + ZERO_WIDTH_SPACE + "upload" in wrapped and "a-" + ZERO_WIDTH_SPACE + "b" in wrapped
    assert breakable_urls(wrapped) == wrapped


def test_parenthetical_tail_of_a_caption_is_not_repeated_in_the_reference_sentence():
    caption = "Рисунок 2 – Визиты доменов с данными Semrush (логарифмическая шкала; в скобках доля)"
    assert figure_reference_sentence(caption) == "Тема «визиты доменов с данными Semrush» показана на рисунке 2."
    assert table_reference_sentence("Таблица 4 – (ЛР5)") == "Данные по теме «(ЛР5)» приведены в таблице 4."
